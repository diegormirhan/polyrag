from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

def load_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")

def load_table(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path)
    return pd.read_excel(path)

def load_document(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path)
    elif suffix == ".docx":
        return _load_docx(path)
    elif suffix == ".pptx":
        return _load_pptx(path)
    raise ValueError(f"unsupported document extension: {suffix}")

def _load_pdf(path: Path) -> str:
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)

_HEADING_STYLE = re.compile(r"^heading (\d)$", re.IGNORECASE)


def _load_docx(path: Path) -> str:
    """Extracts the text, keeping headings as Markdown.

    A .docx already knows which paragraphs are headings — dropping that and
    emitting flat text throws away the only thing that tells a later chunk which
    section it came from. Markdown is the carrier because chunking already
    understands it, and .md files arrive with it for free.
    """
    lines = []
    for paragraph in Document(path).paragraphs:
        level = _HEADING_STYLE.match(paragraph.style.name or "")
        if level and paragraph.text.strip():
            lines.append(f"{'#' * int(level.group(1))} {paragraph.text.strip()}")
        else:
            lines.append(paragraph.text)
    return "\n".join(lines)

def _load_pptx(path: Path) -> str:
    prs = Presentation(path)
    lines: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)
